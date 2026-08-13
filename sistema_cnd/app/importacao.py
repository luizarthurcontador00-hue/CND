"""Importação de empresas em massa por CSV ou Excel.

Aceita o modelo_empresas.csv que acompanha o sistema. É tolerante:
  - reconhece os cabeçalhos com ou sem acento, em maiúscula ou minúscula;
  - aceita CNPJ com ou sem pontuação;
  - aceita ";" ou "," como separador do CSV;
  - aceita arquivos salvos em UTF-8 ou no padrão do Excel brasileiro (latin-1).

Nunca duplica empresa: se o CNPJ já existe, ATUALIZA o cadastro.
"""

from __future__ import annotations

import csv
import io
import logging
import unicodedata
from dataclasses import dataclass, field

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modelos import Empresa, TipoCertidao
from app.servicos import aplicar_exigencias, exigencias_sugeridas
from robos.base import cnpj_valido, limpar_cnpj

logger = logging.getLogger(__name__)

#: Cabeçalho normalizado -> campo do modelo
COLUNAS = {
    "cnpj": "cnpj",
    "razaosocial": "razao_social",
    "razao": "razao_social",
    "nomefantasia": "nome_fantasia",
    "fantasia": "nome_fantasia",
    "uf": "uf",
    "estado": "uf",
    "municipio": "municipio",
    "cidade": "municipio",
    "inscricaoestadual": "inscricao_estadual",
    "ie": "inscricao_estadual",
    "inscricaomunicipal": "inscricao_municipal",
    "im": "inscricao_municipal",
    "regimetributario": "regime_tributario",
    "regime": "regime_tributario",
    "ativa": "ativa",
    "observacoes": "observacoes",
    "certidoes": "certidoes",
    "certidoesexigidas": "certidoes",
}

VERDADEIROS = {"sim", "s", "true", "1", "verdadeiro", "x", "ativa", "ativo"}


@dataclass
class ResultadoImportacao:
    criadas: int = 0
    atualizadas: int = 0
    ignoradas: int = 0
    erros: list[str] = field(default_factory=list)

    @property
    def total_processado(self) -> int:
        return self.criadas + self.atualizadas

    @property
    def resumo(self) -> str:
        partes = [f"{self.criadas} criada(s)", f"{self.atualizadas} atualizada(s)"]
        if self.ignoradas:
            partes.append(f"{self.ignoradas} com problema")
        return ", ".join(partes)


def _normalizar_cabecalho(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", texto or "")
    sem_acento = "".join(c for c in sem_acento if not unicodedata.combining(c))
    return "".join(c for c in sem_acento.lower() if c.isalnum())


def _decodificar(conteudo: bytes) -> str:
    for codificacao in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return conteudo.decode("utf-8", errors="replace")


def _linhas_do_csv(conteudo: bytes) -> list[dict]:
    texto = _decodificar(conteudo)
    amostra = texto[:4096]
    separador = ";" if amostra.count(";") > amostra.count(",") else ","
    leitor = csv.DictReader(io.StringIO(texto), delimiter=separador)
    return [linha for linha in leitor]


def _linhas_do_excel(conteudo: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise ValueError(
            "Para importar .xlsx é preciso instalar o openpyxl "
            "(já está no requirements.txt)."
        ) from e

    planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    aba = planilha.active
    linhas = list(aba.iter_rows(values_only=True))
    if not linhas:
        return []

    cabecalho = [str(c) if c is not None else "" for c in linhas[0]]
    resultado = []
    for linha in linhas[1:]:
        if all(c is None or str(c).strip() == "" for c in linha):
            continue
        resultado.append(
            {cabecalho[i]: linha[i] for i in range(min(len(cabecalho), len(linha)))}
        )
    return resultado


def _mapear(linha: dict) -> dict:
    """Traduz os cabeçalhos da planilha para os campos do modelo."""
    dados: dict = {}
    for chave_bruta, valor in linha.items():
        campo = COLUNAS.get(_normalizar_cabecalho(str(chave_bruta)))
        if not campo:
            continue
        if valor is None:
            continue
        texto = str(valor).strip()
        dados[campo] = texto or None
    return dados


def importar(sessao: Session, conteudo: bytes, nome_arquivo: str) -> ResultadoImportacao:
    """Importa o arquivo e devolve o resumo do que aconteceu."""
    resultado = ResultadoImportacao()

    nome = (nome_arquivo or "").lower()
    try:
        if nome.endswith((".xlsx", ".xlsm")):
            linhas = _linhas_do_excel(conteudo)
        else:
            linhas = _linhas_do_csv(conteudo)
    except Exception as e:
        resultado.erros.append(f"Não consegui ler o arquivo: {e}")
        return resultado

    if not linhas:
        resultado.erros.append("O arquivo está vazio ou sem cabeçalho.")
        return resultado

    tipos_validos = {t.value for t in TipoCertidao}

    for numero, linha_bruta in enumerate(linhas, start=2):  # linha 1 é o cabeçalho
        dados = _mapear(linha_bruta)

        cnpj = limpar_cnpj(dados.get("cnpj") or "")
        if not cnpj:
            resultado.ignoradas += 1
            resultado.erros.append(f"Linha {numero}: sem CNPJ.")
            continue

        if not cnpj_valido(cnpj):
            resultado.ignoradas += 1
            resultado.erros.append(
                f"Linha {numero}: CNPJ {dados.get('cnpj')} é inválido (dígito verificador não confere)."
            )
            continue

        razao = dados.get("razao_social")
        if not razao:
            resultado.ignoradas += 1
            resultado.erros.append(f"Linha {numero}: sem razão social.")
            continue

        certidoes_texto = dados.pop("certidoes", None)
        ativa_texto = dados.pop("ativa", None)

        campos = {
            "razao_social": razao,
            "nome_fantasia": dados.get("nome_fantasia"),
            "uf": (dados.get("uf") or "").upper()[:2] or None,
            "municipio": dados.get("municipio"),
            "inscricao_estadual": dados.get("inscricao_estadual"),
            "inscricao_municipal": dados.get("inscricao_municipal"),
            "regime_tributario": dados.get("regime_tributario"),
            "observacoes": dados.get("observacoes"),
        }
        if ativa_texto is not None:
            campos["ativa"] = ativa_texto.strip().lower() in VERDADEIROS

        empresa = sessao.scalar(select(Empresa).where(Empresa.cnpj == cnpj))
        nova = empresa is None

        if nova:
            empresa = Empresa(cnpj=cnpj, **campos)
            sessao.add(empresa)
            sessao.flush()
        else:
            for campo, valor in campos.items():
                if valor is not None:
                    setattr(empresa, campo, valor)

        # Certidões exigidas: da planilha, ou sugestão automática nas empresas novas.
        if certidoes_texto:
            pedidos = [
                p.strip().upper().replace(" ", "_")
                for p in str(certidoes_texto).replace(";", ",").split(",")
                if p.strip()
            ]
            desconhecidos = [p for p in pedidos if p not in tipos_validos]
            if desconhecidos:
                resultado.erros.append(
                    f"Linha {numero}: certidão(ões) não reconhecida(s): {', '.join(desconhecidos)}. "
                    f"Use: {', '.join(sorted(tipos_validos))}."
                )
            validos = [p for p in pedidos if p in tipos_validos]
            if validos:
                aplicar_exigencias(sessao, empresa, validos)
        elif nova:
            aplicar_exigencias(
                sessao, empresa, [t.value for t in exigencias_sugeridas(empresa)]
            )

        if nova:
            resultado.criadas += 1
        else:
            resultado.atualizadas += 1

    sessao.commit()
    logger.info("Importação concluída: %s", resultado.resumo)
    return resultado
